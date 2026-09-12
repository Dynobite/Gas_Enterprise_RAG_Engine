"""
Engineering Excel Parser for BOM (Bill of Materials), DFMEA Risk Analysis, and Technical Tables.
Provides automatic merged-cell propagation, multi-row header consolidation,
row-level key-value serialization, and rich metadata tagging for precise vector retrieval.
"""

import os
from typing import List, Dict, Any
from src.parsers.base import BaseParser, DocumentChunk

class ExcelParser(BaseParser):
    def can_parse(self, file_path: str) -> bool:
        return file_path.lower().endswith((".xlsx", ".xls"))

    def parse(self, file_path: str) -> List[DocumentChunk]:
        chunks: List[DocumentChunk] = []
        filename = os.path.basename(file_path)

        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)

            doc_type = (
                "Bill of Materials (BOM)" if "BOM" in filename.upper()
                else "DFMEA Risk Analysis" if "DFMEA" in filename.upper()
                else "Engineering Spreadsheet"
            )

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                max_r = ws.max_row
                max_c = ws.max_column
                if max_r < 1:
                    continue

                # 1. Unmerge and populate entire grid (propagates merged headers & cells)
                grid = [[ws.cell(row=r, column=c).value for c in range(1, max_c + 1)] for r in range(1, max_r + 1)]
                for mrange in list(ws.merged_cells.ranges):
                    min_col, min_row, max_col, max_row = mrange.bounds
                    top_left_val = ws.cell(row=min_row, column=min_col).value
                    if top_left_val is not None:
                        for r in range(min_row, max_row + 1):
                            for c in range(min_col, max_col + 1):
                                if r <= max_r and c <= max_c:
                                    grid[r - 1][c - 1] = top_left_val

                # 2. Extract Document / Sheet Metadata header from top rows
                overview_lines = []
                for r_idx in range(min(12, max_r)):
                    items = [str(c).strip().replace('\n', ' ') for c in grid[r_idx] if c is not None and str(c).strip()]
                    if 1 <= len(items) <= 3 and not any(kw in items[0].lower() for kw in ["№", "поз", "обознач", "наименов", "элемент"]):
                        overview_lines.append(" : ".join(items))

                overview_context = " | ".join(overview_lines[:4])

                # 3. Detect Table Header Row(s)
                header_keywords = [
                    "элемент", "функция", "требование", "отказ", "последствия", "причин", "пчр", "rpn",
                    "позиция", "обозначение", "наименование", "материал", "количество", "кол-во", "кол",
                    "код", "нн", "деталь", "чертеж", "масса", "s", "o", "d", "класс"
                ]

                best_header_start = -1
                for r_idx in range(min(25, max_r)):
                    row_items = [str(c).strip().lower() for c in grid[r_idx] if c is not None and str(c).strip()]
                    matches = sum(1 for kw in header_keywords if any(kw in c for c in row_items))
                    if matches >= 2 and len(row_items) >= 2:
                        best_header_start = r_idx
                        break

                if best_header_start == -1:
                    best_header_start = 0

                # Check if next row is a multi-row sub-header (common in DFMEA and BOM tables)
                has_sub_header = False
                if best_header_start + 1 < max_r:
                    sub_row = grid[best_header_start + 1]
                    sub_matches = sum(
                        1 for kw in header_keywords + ["обнаружение", "предупреждение", "тяжесть", "вероятность", "масса", "примечание", "действия", "статус"]
                        if any(kw in str(c).strip().lower() for c in sub_row if c is not None)
                    )
                    if sub_matches >= 2:
                        has_sub_header = True

                # Consolidate column headers
                headers = []
                for c in range(max_c):
                    h1 = str(grid[best_header_start][c]).strip().replace('\n', ' ') if grid[best_header_start][c] is not None else ""
                    if has_sub_header:
                        h2 = str(grid[best_header_start + 1][c]).strip().replace('\n', ' ') if grid[best_header_start + 1][c] is not None else ""
                        if h1 and h2 and h1.lower() != h2.lower():
                            headers.append(f"{h1} [{h2}]")
                        elif h2:
                            headers.append(h2)
                        elif h1:
                            headers.append(h1)
                        else:
                            headers.append(f"Колонка_{c+1}")
                    else:
                        headers.append(h1 if h1 else f"Колонка_{c+1}")

                data_start_idx = best_header_start + (2 if has_sub_header else 1)

                # 4. Serialize each data row with exact column headers
                serialized_rows = []
                for row_num in range(data_start_idx, max_r):
                    row = grid[row_num]
                    if not any(c is not None and str(c).strip() for c in row):
                        continue

                    row_parts = []
                    for col_idx, cell_val in enumerate(row):
                        if col_idx < len(headers) and cell_val is not None and str(cell_val).strip():
                            header_name = headers[col_idx]
                            val_str = str(cell_val).strip().replace('\n', ' ')
                            row_parts.append(f"{header_name}: {val_str}")

                    if row_parts:
                        serialized_rows.append(f"[Строка {row_num + 1}] " + " | ".join(row_parts))

                # 5. Group serialized rows into semantic micro-chunks
                if serialized_rows:
                    batch_size = 4  # 4 rows per chunk gives rich context + high precision
                    total_batches = (len(serialized_rows) + batch_size - 1) // batch_size
                    for b_idx in range(0, len(serialized_rows), batch_size):
                        batch = serialized_rows[b_idx:b_idx + batch_size]

                        chunk_header = f"=== {doc_type.upper()} ===\nДокумент: {filename}\nЛист: '{sheet_name}'"
                        if overview_context:
                            chunk_header += f"\nКонтекст листа: {overview_context}"

                        chunk_text = f"{chunk_header}\n\n" + "\n".join(batch)

                        metadata = {
                            "source": filename,
                            "file_path": os.path.abspath(file_path),
                            "sheet": sheet_name,
                            "page": (b_idx // batch_size) + 1,
                            "total_pages": total_batches,
                            "doc_type": doc_type,
                            "format": "xlsx",
                            "verification_link": f"http://192.168.152.38:8000/api/documents/{filename}"
                        }
                        chunks.append(DocumentChunk(text=chunk_text, metadata=metadata))

        except Exception as e:
            print(f"Error parsing Excel file {file_path}: {e}")

        return chunks
