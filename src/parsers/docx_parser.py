"""
DOCX Parser supporting technical defect acts, inspection reports, and Word tables.
"""

import os
import re
from typing import List, Dict, Any
from src.parsers.base import BaseParser, DocumentChunk

class DocxParser(BaseParser):
    def can_parse(self, file_path: str) -> bool:
        return file_path.lower().endswith(".docx")

    def parse(self, file_path: str) -> List[DocumentChunk]:
        chunks = []
        filename = os.path.basename(file_path)
        
        try:
            import docx
            doc = docx.Document(file_path)
            
            doc_type = "Defect Act / Protocol" if "АКТ" in filename.upper() or "ДЕФЕКТ" in filename.upper() else "Technical Word Document"
            
            # 1. Process paragraphs
            current_section = "Введение"
            paragraph_texts = []
            
            for p in doc.paragraphs:
                txt = p.text.strip()
                if not txt:
                    continue
                if p.style.name.startswith("Heading") or len(txt) < 80 and txt.isupper():
                    current_section = txt
                paragraph_texts.append(f"[{current_section}] {txt}")

            if paragraph_texts:
                combined_paragraphs = "\n".join(paragraph_texts)
                metadata = {
                    "source": filename,
                    "file_path": os.path.abspath(file_path),
                    "page": 1,
                    "total_pages": 1,
                    "doc_type": doc_type,
                    "content_type": "paragraphs",
                    "format": "docx",
                    "verification_link": f"file://{os.path.abspath(file_path).replace(os.sep, '/')}"
                }
                chunks.append(DocumentChunk(text=combined_paragraphs, metadata=metadata))

            # 2. Process tables
            for t_idx, table in enumerate(doc.tables, start=1):
                rows_text = []
                headers = []
                
                # Extract headers if available
                if len(table.rows) > 0:
                    headers = [cell.text.strip().replace('\n', ' ') for cell in table.rows[0].cells]
                
                for row_idx, row in enumerate(table.rows[1:], start=2):
                    cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    if any(cells):
                        if headers and len(headers) == len(cells):
                            row_desc = " | ".join(f"{h}: {c}" for h, c in zip(headers, cells) if c)
                        else:
                            row_desc = " | ".join(c for c in cells if c)
                        rows_text.append(f"Строка {row_idx}: {row_desc}")

                if rows_text:
                    table_content = f"Таблица #{t_idx} из документа '{filename}':\n" + "\n".join(rows_text)
                    metadata = {
                        "source": filename,
                        "file_path": os.path.abspath(file_path),
                        "page": 1,
                        "total_pages": 1,
                        "doc_type": doc_type,
                        "content_type": f"table_{t_idx}",
                        "format": "docx",
                        "verification_link": f"file://{os.path.abspath(file_path).replace(os.sep, '/')}"
                    }
                    chunks.append(DocumentChunk(text=table_content, metadata=metadata))

        except Exception as e:
            print(f"Error parsing DOCX {file_path}: {e}")

        return chunks
